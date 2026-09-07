import os
import io
from fastapi import FastAPI, HTTPException, Request, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.calculator import CalculoEntrada, CalculoResultado, calcular_presupuesto
from app.database import (
    init_db,
    get_all_pedidos,
    get_pedido,
    create_pedido,
    update_pedido,
    delete_pedido,
    get_gastos_material,
    add_gasto_material,
    delete_gasto_material,
    get_gastos_maquinaria,
    add_gasto_maquinaria,
    delete_gasto_maquinaria,
    get_all_stock,
    update_stock_item,
    descontar_stock,
    get_all_fornituras,
    update_fornitura,
    get_balance_financiero
)

app = FastAPI(title="W&P - 3D Printing & Workshop Manager", version="1.0.0")

# Setup templates and static
BASE_DIR = os.path.dirname(__file__)
app.mount("/static", StaticFiles(directory=os.path.join(BASE_DIR, "static")), name="static")
templates = Jinja2Templates(directory=os.path.join(BASE_DIR, "templates"))

@app.on_event("startup")
def on_startup():
    init_db()

@app.get("/")
def home(request: Request):
    return templates.TemplateResponse(request=request, name="index.html")

# --- API CALCULADORA ---
@app.post("/api/calcular", response_model=CalculoResultado)
def api_calcular(datos: CalculoEntrada):
    try:
        return calcular_presupuesto(datos)
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

# --- API PEDIDOS ---
@app.get("/api/pedidos")
def api_get_pedidos(filtro: str = None, q: str = None):
    return get_all_pedidos(filtro_estado=filtro, query=q)

@app.get("/api/pedidos/{pedido_id}")
def api_get_pedido(pedido_id: str):
    p = get_pedido(pedido_id)
    if not p:
        raise HTTPException(status_code=404, detail="Pedido no encontrado")
    return p

@app.post("/api/pedidos")
def api_create_pedido(datos: dict):
    return create_pedido(datos)

@app.put("/api/pedidos/{pedido_id}")
def api_update_pedido(pedido_id: str, datos: dict):
    p = update_pedido(pedido_id, datos)
    if not p:
        raise HTTPException(status_code=404, detail="Pedido no encontrado")
    return p

@app.delete("/api/pedidos/{pedido_id}")
def api_delete_pedido(pedido_id: str):
    ok = delete_pedido(pedido_id)
    if not ok:
        raise HTTPException(status_code=404, detail="Pedido no encontrado")
    return {"status": "ok", "deleted": pedido_id}

# --- API STOCK ---
@app.get("/api/stock")
def api_get_stock():
    return get_all_stock()

@app.put("/api/stock/{stock_id}")
def api_update_stock(stock_id: int, datos: dict):
    s = update_stock_item(stock_id, datos)
    if not s:
        raise HTTPException(status_code=404, detail="Material no encontrado")
    return s

@app.post("/api/stock/descontar")
def api_descontar_stock(datos: dict):
    stock_id = int(datos.get("stock_id", 0))
    gramos = float(datos.get("gramos", 0.0))
    socio = str(datos.get("socio", "J"))
    res = descontar_stock(stock_id, gramos, socio)
    if not res:
        raise HTTPException(status_code=404, detail="Material no encontrado")
    return res

# --- API FORNITURAS ---
@app.get("/api/fornituras")
def api_get_fornituras():
    return get_all_fornituras()

@app.put("/api/fornituras/{fornitura_id}")
def api_update_fornitura(fornitura_id: int, datos: dict):
    cantidad = float(datos.get("cantidad", 0.0))
    f = update_fornitura(fornitura_id, cantidad)
    if not f:
        raise HTTPException(status_code=404, detail="Fornitura no encontrada")
    return f

# --- API GASTOS ---
@app.get("/api/gastos/material")
def api_get_gastos_material():
    return get_gastos_material()

@app.post("/api/gastos/material")
def api_add_gasto_material(datos: dict):
    return add_gasto_material(datos)

@app.delete("/api/gastos/material/{gasto_id}")
def api_delete_gasto_material(gasto_id: int):
    ok = delete_gasto_material(gasto_id)
    if not ok:
        raise HTTPException(status_code=404, detail="Gasto no encontrado")
    return {"status": "ok"}

@app.get("/api/gastos/maquinaria")
def api_get_gastos_maquinaria():
    return get_gastos_maquinaria()

@app.post("/api/gastos/maquinaria")
def api_add_gasto_maquinaria(datos: dict):
    return add_gasto_maquinaria(datos)

@app.delete("/api/gastos/maquinaria/{gasto_id}")
def api_delete_gasto_maquinaria(gasto_id: int):
    ok = delete_gasto_maquinaria(gasto_id)
    if not ok:
        raise HTTPException(status_code=404, detail="Gasto no encontrado")
    return {"status": "ok"}

# --- API BALANCE ---
@app.get("/api/balance")
def api_get_balance():
    return get_balance_financiero()

# --- EXPORTAR EXCEL ---
@app.get("/api/export/excel")
def api_export_excel():
    wb = openpyxl.Workbook()
    
    # 1. Pestaña Pedidos
    ws_pedidos = wb.active
    ws_pedidos.title = "Pedidos"
    headers_pedidos = ["Nº Pedido", "Título", "Cliente", "Fecha", "Unidades", "Precio Ud", "Precio Total", "% CI", "Gastos Indirectos", "Beneficio", "¿Quién Hace?", "¿Quién Cobra?"]
    ws_pedidos.append(headers_pedidos)
    
    for cell in ws_pedidos[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        
    pedidos = get_all_pedidos()
    for p in reversed(pedidos):
        ws_pedidos.append([
            p["id"], p["titulo"], p["cliente"], p["fecha"],
            p["unidades"], p["precio_unidad"], p["precio_total"],
            p["pct_ci"], p["gastos_indirectos"], p["beneficio"],
            p["quien_hace"], p["quien_cobra"]
        ])
        
    # 2. Pestaña Gastos Material
    ws_mat = wb.create_sheet(title="Gastos Material")
    ws_mat.append(["ID", "Concepto", "Fecha", "Precio (€)", "Pagado Por", "Enlace"])
    for cell in ws_mat[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    for m in get_gastos_material():
        ws_mat.append([m["id"], m["concepto"], m["fecha"], m["precio"], m["pagado_por"], m["enlace"]])
        
    # 3. Pestaña Stock
    ws_stock = wb.create_sheet(title="Stock Materiales")
    ws_stock.append(["Tipo", "Color", "Stock Total (g)", "Pablo (g)", "Javi (g)", "Precio kg (€)"])
    for cell in ws_stock[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4338CA", end_color="4338CA", fill_type="solid")
    for s in get_all_stock():
        ws_stock.append([s["tipo"], s["color"], s["stock_total_g"], s["pablo_g"], s["javi_g"], s["precio_kg_estimado"]])
        
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=W_and_P_Actualizado.xlsx"}
    )
